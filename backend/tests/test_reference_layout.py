from io import BytesIO
from zipfile import ZipFile
from datetime import date
from types import SimpleNamespace
from PIL import Image
from docx import Document
from app.schemas import ReportData
from app.calculations import calculate, review_issues
from app.reporting import docx_bytes, TEMPLATE
from app.reference_layout import layout_context, dms, money_reference
from app.config import settings
from tests.test_workflow import complete_data


def xml(raw):
    with ZipFile(BytesIO(raw)) as z:return z.read('word/document.xml').decode()


def test_reference_layout_is_dynamic_and_safe_for_another_property():
    data=ReportData(**complete_data())
    data.assignment.display_addressee='Separate recipient <A & B>'
    data.assignment.display_address='New recipient address'
    data.valuer.email='different@example.test'
    data.findings.documents='First document\nSecond document'
    data.findings.limitations='First limitation.\n\nSecond limitation.'
    rendered=docx_bytes(data,'draft',[])
    doc=Document(BytesIO(rendered));text='\n'.join(p.text for p in doc.paragraphs)
    assert 'Separate recipient <A & B>' in text
    assert 'First document' in text and 'Second document' in text
    assert len([p for p in doc.paragraphs if p.text.startswith('* ')])==2
    assert 'DRAFT' in text
    assert 'Roof' not in '\n'.join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    assert 'valuation-reference-v2' in str(TEMPLATE)
    assert '{{' not in xml(rendered) and '{%' not in xml(rendered)
    assert all('different@example.test' in '\n'.join(p.text for p in s.footer.paragraphs) for s in doc.sections)


def test_mixed_floor_rates_and_adjustments_are_printed():
    data=ReportData(**complete_data());data.property.kind='land_building'
    raw=data.model_dump(mode='json');raw['building']['floors']=[
        {'name':'Floor A','area_sqft':'100','rate_per_sqft':'2000'},
        {'name':'Floor B','area_sqft':'50','rate_per_sqft':'3000'}]
    raw['valuation'].update(site_improvements='1000',depreciation_amount='500',adjustment_amount='-250',adjustment_reason='Test adjustment')
    data=ReportData(**raw);context=layout_context(data,calculate(data))
    assert 'Floor A' in str(context['building_calculation'])
    assert '2,000/-' in str(context['building_calculation']) and '3,000/-' in str(context['building_calculation'])
    text=xml(docx_bytes(data,'draft',[]))
    assert '350,000/-' in text and 'Test adjustment' in text
    assert 'Separate site improvements' in text and 'Less building depreciation' in text


def test_photo_roles_generate_all_requested_slots():
    folder=settings.data_dir/'photos';folder.mkdir(exist_ok=True)
    assets=[]
    for i,role in enumerate(['cover','access','interior','interior','interior']):
        filename=f'role-test-{i}.png';Image.new('RGB',(80+i,60),'white').save(folder/filename)
        assets.append(SimpleNamespace(filename=filename,caption=f'[{role}] Caption {i}'))
    raw=complete_data();raw['property']['kind']='land_building'
    raw['building']['floors']=[{'name':'Floor A','area_sqft':'100','rate_per_sqft':'1000'}]
    doc=Document(BytesIO(docx_bytes(ReportData(**raw),'draft',assets)))
    assert len(doc.inline_shapes)==5
    assert 'Caption 4' in '\n'.join(p.text for p in doc.paragraphs)


def test_missing_photos_never_reuse_source_property_images():
    with ZipFile(TEMPLATE) as z:
        assert not any(n.lower().endswith(('.jpg','.jpeg')) for n in z.namelist())
    data=ReportData(**complete_data())
    doc=Document(BytesIO(docx_bytes(data,'draft',[])))
    assert len(doc.inline_shapes)==0


def test_dates_coordinates_and_optional_values():
    assert dms(0,True)=='0°00\'00.0"N'
    assert dms(-1.5,False)=='1°30\'00.0"W'
    assert money_reference('100.25')=='100.25'
    data=ReportData(**complete_data());data.assignment.valuation_date=date(2027,1,1)
    assert 'Valuation date cannot be after the report date.' in review_issues(data)
    out=xml(docx_bytes(data,'draft',[]))
    assert 'Forced sale Value of the subject property' not in out
    assert 'Insurance sale Value of the subject property' not in out
