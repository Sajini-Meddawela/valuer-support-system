from copy import deepcopy
from pathlib import Path
from io import BytesIO
from zipfile import ZipFile, ZIP_DEFLATED
import argparse
from docx import Document
from docx.shared import Pt, RGBColor, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from lxml import etree

OUT = Path(__file__).resolve().parents[1] / 'templates' / 'valuation-reference-v2.docx'


def put(p, text, keep_art=False):
    """Preserve paragraph geometry and the source's first text-run formatting."""
    old = next((r for r in p.runs if r.text.strip()), None)
    style = deepcopy(old._r.rPr) if old is not None and old._r.rPr is not None else None
    if keep_art:
        for node in p._p.xpath('.//w:t'):
            node.text = ''
    else:
        for child in list(p._p):
            if child.tag != qn('w:pPr'):
                p._p.remove(child)
    run = p.add_run(text)
    if style is not None:
        run._r.insert(0, style)
    return run


def fragments(p, values):
    put(p, '')
    for text, color, bold in values:
        r = p.add_run(text)
        r.font.name = 'Trebuchet MS'; r.font.size = Pt(11)
        r.font.color.rgb = RGBColor.from_string(color)
        r.bold = bold


def control(anchor, text, before=True):
    node = OxmlElement('w:p')
    if before: anchor.addprevious(node)
    else: anchor.addnext(node)
    p = Paragraph(node, None)
    p.add_run(text).font.size = Pt(1)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1
    return p


def wrap(first, last, expr):
    control(first, '{%p if ' + expr + ' %}')
    control(last, '{%p endif %}', False)


def create(source, output=OUT):
    d = Document(source); ps = d.paragraphs; ts = d.tables
    if len(ps) != 118 or len(ts) != 3 or ps[38].text.strip() != 'Location Sketch (Not to Scale/ Source: Google Map)':
        raise ValueError('Reference layout differs from the supplied version. No template was written.')
    for i in range(5):
        put(ps[i], '{{ letterhead[' + str(i) + '] }}')
    for i, token in [(6,'{{ valuer.address }}'),(7,'{{ phones[0] }}'),(8,'{{ phones[1] }}'),(9,'{{ phones[2] }}'),(10,'{{ valuer.email }}')]:
        put(ps[i],token,keep_art=i in (6,7))
    fragments(ps[13],[('My Ref: ','000000',True),('{{ assignment.reference }}','FF0000',True),
        ('\tYour Ref: {{ assignment.bank_reference }}\tDate:{{ assignment.report_date|shortdate }}','000000',True)])
    for r in ps[13].runs:r.font.size=Pt(12)
    ps[13].paragraph_format.tab_stops.clear_all()
    ps[13].paragraph_format.tab_stops.add_tab_stop(Mm(65))
    ps[13].paragraph_format.tab_stops.add_tab_stop(Mm(119))
    put(ps[14], '(Private and Confidential){% if not is_final %} — DRAFT{% endif %}')
    put(ps[16], 'Valuation Report of Land bearing Lot No. {{ property.lot_number }}, in Survey Plan No. {{ property.plan_number }}, dated {{ property.plan_date|longdate }} , made by {{ property.surveyor }}')
    put(ps[17], '{{ cover_image }}')
    put(ps[22], '{{ display_addressee }}')
    put(ps[23], '{{ display_address }}')
    fragments(ps[26],[('Following your recent request on ','000000',True),('{{ assignment.request_date|longdate }}','FF0000',True),
        (' to value the above property {{ purpose_clause }}, I have inspected the aforesaid property on ','000000',True),
        ('{{ assignment.inspection_date|longdate }}','FF0000',True),(' and property has been valued on behalf of the ','000000',True),
        ('{{ bank.name }}','FF0000',True),(' and not for any other person & no responsibility is accepted to third parties for the whole or any part of the contents.','000000',True)])
    mapping={29:'{{r situation_rich }}',31:'{{r identity_rich }}',
        34:'{{ property.access }}',37:'GPS coordination for the property: {{ coordinates_dms }}',
        39:'{{ map_image }}',42:'{{ access_image }} {{ qr_image }}',43:'Note: {{ property.road_width }}',
        45:'{{ property.locality }}',46:'{{ property.services }}',49:'{{ property.land_description }}',
        50:'{{ boundary_note }}',51:'North by\t: {{ property.north }}',52:'East by\t: {{ property.east }}',
        53:'South by\t: {{ property.south }}',54:'West by\t: {{ property.west }}',
        55:'Description of the property – {{ building_name }}',58:'{{ interior_left }} {{ interior_right }}',
        60:'{{r assessment_rich }}',63:'{{ document_line }}',64:'{{ findings.approvals }}',
        67:'{{ findings.street_building_lines }}',69:'{{ findings.legal_interest }}',
        70:'{{r ownership_rich }}',72:'{{ deed_text }}',75:'{{ findings.analysis }}',78:'{{r evidence_rich }}',
        83:'The {{ findings.method }} uses as the approach to ascertain the present market value in this valuation. The subject land is considered the {{ extent_hyphen }} in extent as per the survey plan.',
        87:'{{r land_calculation }}',90:'{{r building_calculation }}',92:'{{r total_calculation }}',93:'{{r rounded_calculation }}',
        96:'{{r market_summary }}',97:'{{r forced_summary }}',98:'{{r insurance_summary }}',
        101:'{{r limitation }}',103:'',105:'',107:'',109:'',111:'{{ findings.certification }}',
        115:'{{ valuer.name }}',116:'{{ signature_designation }}',117:'{{ assignment.report_date|shortdate }}'}
    for i,text in mapping.items():put(ps[i],text)
    control(ps[63]._p, '{%p for document_line in document_lines %}')
    control(ps[63]._p, '{%p endfor %}', False)
    control(ps[101]._p, '{%p for limitation in limitation_paragraphs %}')
    control(ps[101]._p, '{%p endfor %}', False)
    ps[101].paragraph_format.space_after = Pt(12)
    for i in range(102,110):
        ps[i]._p.getparent().remove(ps[i]._p)
    for i,condition in [(50,'boundary_note'),(64,'findings.approvals'),(70,'ownership_text'),(72,'deed_text'),
                        (97,'valuation.forced_sale_value is not none'),(98,'valuation.insurance_value is not none')]:
        wrap(ps[i]._p,ps[i]._p,condition)
    for i in [28,30,33,38,41,44,48,55,59,62,66,68,74,77,82,85,86,89,95,99,110,114,115,116]:
        ps[i].paragraph_format.keep_with_next=True
    ps[111].paragraph_format.keep_with_next=True
    for i in [112,113]:ps[i].paragraph_format.keep_with_next=True
    cells={(0,1,0):'Applicant : {{ assignment.applicant }}',(0,2,0):'Address  : {{ assignment.applicant_address }}',
       (0,3,0):'Present Owner:    {{ assignment.owner }}',(0,4,0):'Survey Plan No  : {{ property.plan_number }}',
       (0,4,1):'Lot No\t : {{ property.lot_number }}',(0,5,0):'Date                  : {{ property.plan_date|longdate }}',
       (0,5,1):'Extent         : {{ extent_spaces }}',(0,6,0):'Made By : {{ property.surveyor }}',
       (0,6,1):'Village         : {{ property.village }}',(0,7,0):'Local Authority    : {{ property.local_authority }}',
       (1,0,0):'{{ building.description }}',(2,1,0):'{{ building_name }}',(2,1,1):'{{ building.area_source }}',(2,1,2):'{{ floor_area_lines }}'}
    for row,key in enumerate(['roof','walls','floor_finish','doors_windows','conveniences','condition','occupier'],1):
        cells[(1,row,1)]='{{ building.'+key+' }}'
    for (ti,ri,ci),value in cells.items():
        cell=ts[ti].cell(ri,ci);put(cell.paragraphs[0],value)
        for p in cell.paragraphs[1:]:p._p.getparent().remove(p._p)
    ts[2].cell(1,2).paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.LEFT
    ps[63].alignment=WD_ALIGN_PARAGRAPH.LEFT
    if ps[64]._p.pPr is not None:
        for el in list(ps[64]._p.pPr):
            if el.tag==qn('w:numPr'):ps[64]._p.pPr.remove(el)
    wrap(ps[55]._p,ps[58]._p,"property.kind == 'land_building'")
    wrap(ps[89]._p,ps[90]._p,"property.kind == 'land_building'")
    extra=control(ps[92]._p,'{%p if extra_financial_lines %}')
    e=control(ps[92]._p,'{{ extra_financial_lines }}');e.paragraph_format.line_spacing=1.5
    e.runs[0].font.size=Pt(11);e.runs[0].font.name='Trebuchet MS'
    control(ps[92]._p,'{%p endif %}')
    extra=control(ps[59]._p,'{%p for photo in extra_photos %}')
    control(ps[59]._p,'{{ photo.image }}').runs[0].font.size=Pt(11)
    control(ps[59]._p,'{{ photo.caption }}').runs[0].font.size=Pt(10)
    control(ps[59]._p,'{%p endfor %}')
    for section in d.sections:
        for p in section.footer.paragraphs:
            for r in p.runs:
                if '@' in r.text:r.text='{{ valuer.email }}'
    refs=set(d.element.xpath('//@r:embed | //@r:id | //@r:link'))
    for rid,rel in list(d.part.rels.items()):
        if rel.reltype.endswith('/image') and rid not in refs:del d.part.rels[rid]
        elif rel.is_external:del d.part.rels[rid]
    for node in d.element.xpath('.//wp:docPr | .//pic:cNvPr'):
        node.set('name','Template graphic')
        for key in ['descr','title']:
            node.attrib.pop(key,None)
    cp=d.core_properties
    for key in ['author','last_modified_by','subject','keywords','comments','category','identifier']:
        setattr(cp,key,'')
    cp.title='Valuation report template';cp.revision=1
    buf=BytesIO();d.save(buf)
    ns={'r':'http://schemas.openxmlformats.org/package/2006/relationships'}
    with ZipFile(buf) as zin, ZipFile(output,'w',ZIP_DEFLATED) as zout:
        removed={n for n in zin.namelist() if n.startswith('customXml/') or n.startswith('docProps/thumbnail') or n=='docProps/custom.xml'}
        for name in zin.namelist():
            if name in removed:continue
            raw=zin.read(name)
            if name.endswith('.rels'):
                tree=etree.fromstring(raw)
                for rel in list(tree):
                    target=rel.get('Target','')
                    if 'customXml' in target or 'thumbnail' in target or target.endswith('custom.xml'):tree.remove(rel)
                raw=etree.tostring(tree,xml_declaration=True,encoding='UTF-8',standalone=True)
            elif name=='[Content_Types].xml':
                tree=etree.fromstring(raw)
                for child in list(tree):
                    if child.get('PartName','').lstrip('/') in removed:tree.remove(child)
                raw=etree.tostring(tree,xml_declaration=True,encoding='UTF-8',standalone=True)
            zout.writestr(name,raw)
    print(f'Created sanitised template: {output}')


if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('source',type=Path)
    parser.add_argument('--output',type=Path,default=OUT)
    args=parser.parse_args();args.output.parent.mkdir(parents=True,exist_ok=True)
    create(args.source,args.output)
