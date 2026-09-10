"""Create a clean, editable template. No content from the confidential sample is bundled."""
from pathlib import Path
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'templates' / 'valuation-v1.docx'


def create():
    d = Document()
    s = d.sections[0]
    s.page_width = Mm(210); s.page_height = Mm(297)
    s.top_margin = Mm(20); s.bottom_margin = Mm(20)
    s.left_margin = Mm(25); s.right_margin = Mm(25)
    s.header_distance = Mm(8); s.footer_distance = Mm(9)
    for name in ['Normal', 'Body Text']:
        style = d.styles[name]
        style.font.name = 'Liberation Serif'; style.font.size = Pt(11)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.08
    for name, size in [('Title', 20), ('Heading 1', 14), ('Heading 2', 12)]:
        st = d.styles[name]; st.font.name = 'Liberation Serif'; st.font.size = Pt(size)
        st.font.bold = True; st.font.color.rgb = RGBColor.from_string('182C3D')
        st.paragraph_format.space_before = Pt(12); st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.keep_with_next = True
    header = s.header.paragraphs[0]
    header.text = '{{ assignment.reference }}  |  PRIVATE AND CONFIDENTIAL'
    header.style = d.styles['Normal']; header.runs[0].font.size = Pt(9)
    footer = s.footer.paragraphs[0]
    footer.text = '{{ status }}  |  Page '
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
    for r in footer.runs: r.font.size = Pt(9)

    def p(text):
        d.add_paragraph(text)

    def heading(text):
        d.add_heading(text, 1)

    def table(rows, widths=(45, 115)):
        t = d.add_table(rows=0, cols=len(widths)); t.style = 'Table Grid'; t.autofit = False
        for col, width in zip(t.columns, widths): col.width = Mm(width)
        tblPr = t._tbl.tblPr
        tblW = tblPr.find(qn('w:tblW')); tblW.set(qn('w:type'), 'dxa'); tblW.set(qn('w:w'), str(round(sum(widths) / 25.4 * 1440)))
        for values in rows:
            cells = t.add_row().cells
            for c, value, width in zip(cells, values, widths):
                c.width = Mm(width); c.text = value
                for para in c.paragraphs: para.paragraph_format.space_after = Pt(4)
            cant = OxmlElement('w:cantSplit'); t.rows[-1]._tr.get_or_add_trPr().append(cant)
        return t

    d.add_heading('{{ valuer.name|blank }}', 0)
    p('{{ valuer.qualifications|blank }}\n{{ valuer.registration|blank }}\n{{ valuer.address|blank }}\n{{ valuer.phone|blank }}  |  {{ valuer.email|blank }}')
    d.add_heading('Valuation Report', 1)
    p('{{ status }}')
    table([
        ('Our reference', '{{ assignment.reference|blank }}'),
        ('Bank reference', '{{ assignment.bank_reference|blank }}'),
        ('Report date', '{{ assignment.report_date|blank }}'),
        ('Prepared for', '{{ bank.name|blank }}\n{{ bank.branch|blank }}\n{{ bank.recipient|blank }}\n{{ bank.address|blank }}'),
        ('Purpose', '{{ assignment.purpose|blank }}'),
        ('Request / inspection / valuation dates', '{{ assignment.request_date|blank }} / {{ assignment.inspection_date|blank }} / {{ assignment.valuation_date|blank }}'),
        ('Applicant', '{{ assignment.applicant|blank }}\n{{ assignment.applicant_address|blank }}'),
        ('Present owner', '{{ assignment.owner|blank }}\n{{ assignment.owner_address|blank }}'),
        ('Property', '{{ property.address|blank }}'),
    ])
    heading('Synopsis of the valuation')
    table([
        ('Land / lot', '{{ property.land_name|blank }} / {{ property.lot_number|blank }}'),
        ('Survey plan', '{{ property.plan_number|blank }} dated {{ property.plan_date|blank }}'),
        ('Surveyor', '{{ property.surveyor|blank }}'),
        ('Extent', '{{ property.acres }}A–{{ property.roods }}R–{{ property.perches }}P ({{ calc.extent_perches }} perches)'),
        ('Village / district / province', '{{ property.village|blank }} / {{ property.district|blank }} / {{ property.province|blank }}'),
        ('Local authority', '{{ property.local_authority|blank }}'),
        ('Assessment number', '{{ property.assessment_number|blank }}'),
    ])
    d.add_page_break()
    heading('Situation, identification and access')
    p('{{ property.situation|blank }}')
    p('The property is identified by survey plan {{ property.plan_number|blank }}, lot {{ property.lot_number|blank }}, prepared by {{ property.surveyor|blank }}.')
    p('{{ property.access|blank }}')
    p('Access road width: {{ property.road_width|blank }}')
    p('Coordinates: {{ property.latitude|blank }}, {{ property.longitude|blank }}')
    heading('Location map and access')
    p('{{ map_image }}'); p('{{ map_note }}')
    p('{%p if property.latitude is not none %}')
    p('{{ qr_image }}'); p('Scan to open the supplied location in Google Maps.')
    p('{%p endif %}')
    heading('Property and locality')
    p('{{ property.locality|blank }}')
    p('Services and infrastructure: {{ property.services|blank }}')
    heading('Land description and boundaries')
    p('{{ property.land_description|blank }}')
    table([(label, '{{ property.' + key + '|blank }}') for label, key in [('North','north'),('East','east'),('South','south'),('West','west')]])
    p('{%p if property.kind == "land_building" %}')
    d.add_page_break()
    heading('Building description')
    p('{{ building.description|blank }}')
    table([(label, '{{ building.' + key + '|blank }}') for label, key in [
        ('Roof','roof'),('Walls','walls'),('Floor finish','floor_finish'),('Doors and windows','doors_windows'),
        ('Conveniences','conveniences'),('Condition','condition'),('Occupier','occupier'),('Floor-area source','area_source')]])
    t = table([('Floor', 'Area (sq ft)', 'Rate (LKR / sq ft)')], (65,45,50))
    for row in [('{%tr for floor in building.floors %}', '', ''), ('{{ floor.name }}','{{ floor.area_sqft }}','{{ floor.rate_per_sqft|money }}'),('{%tr endfor %}','','')]:
        cells = t.add_row().cells
        for c, text in zip(cells,row): c.text = text
    p('{%p endif %}')
    d.add_page_break()
    heading('Documents, approvals and legal information')
    for label, key in [('Documents reviewed','documents'),('Approvals','approvals'),('Street lines and building limits','street_building_lines'),('Legal interest and assumptions','legal_interest')]:
        d.add_heading(label,2); p('{{ findings.'+key+'|blank }}')
    p('Deed: {{ findings.deed_number|blank }}\nDeed date: {{ findings.deed_date|blank }}\nNotary: {{ findings.notary|blank }}')
    heading('Basis and analysis of valuation')
    p('Adopted method: {{ findings.method|blank }}')
    d.add_heading('Comparable evidence',2); p('{{ findings.comparable_evidence|blank }}')
    d.add_heading('Analysis and adopted rates',2); p('{{ findings.analysis|blank }}')
    heading('Valuation calculations')
    table([
        ('Land', '{{ calc.extent_perches }} perches × LKR {{ valuation.land_rate_per_perch|money }} = LKR {{ calc.land_value|money }}'),
        ('Gross building value', 'LKR {{ calc.building_value|money }}'),
        ('Site improvements', 'LKR {{ valuation.site_improvements|money }}'),
        ('Building depreciation', 'LKR {{ valuation.depreciation_amount|money }}'),
        ('Other adjustment', 'LKR {{ valuation.adjustment_amount|money }}\n{{ valuation.adjustment_reason|blank }}'),
        ('Calculated total', 'LKR {{ calc.calculated_total|money }}'),
        ('Rounded indication', 'LKR {{ calc.rounded_value|money }}'),
        ('Adopted market value', 'LKR {{ calc.market_value|money }}\n{{ calc.market_value_words }}'),
    ])
    p('{%p if valuation.final_market_value is not none %}')
    p('Reason for adopted market value: {{ valuation.final_market_reason|blank }}')
    p('{%p endif %}')
    p('{%p if valuation.forced_sale_value is not none %}')
    p('Forced-sale value: LKR {{ valuation.forced_sale_value|money }} ({{ valuation.forced_sale_value|words }}).\nBasis: {{ valuation.forced_sale_reason|blank }}')
    p('{%p endif %}')
    p('{%p if valuation.insurance_value is not none %}')
    p('Insurance value: LKR {{ valuation.insurance_value|money }} ({{ valuation.insurance_value|words }}).\nBasis: {{ valuation.insurance_reason|blank }}')
    p('{%p endif %}')
    heading('Limitations and assumptions')
    p('{{ findings.limitations|blank }}')
    p('Intended recipient: {{ bank.name|blank }}. Stated purpose: {{ assignment.purpose|blank }}.')
    heading('Certification')
    p('{{ findings.certification|blank }}')
    d.paragraphs[-1].paragraph_format.keep_with_next = True
    p('Signature: __________________________\n{{ valuer.name|blank }}\n{{ valuer.registration|blank }}\n{{ assignment.report_date|blank }}')
    p('Template: {{ template_version }}')
    p('{%p if photos %}')
    d.add_page_break(); heading('Inspection photographs')
    p('{%p for photo in photos %}'); p('{{ photo.image }}'); p('{{ photo.caption }}'); p('{%p endfor %}')
    p('{%p endif %}')
    d.core_properties.author = 'Valuer Support'
    d.core_properties.title = 'Clean valuation report template'
    # Keep control-only rows/paragraphs compact in the editable source template.
    # docxtpl removes these control containers in the rendered report.
    controls = list(d.paragraphs)
    for t in d.tables:
        for row in t.rows:
            for cell in row.cells:
                controls.extend(cell.paragraphs)
    for para in controls:
        if para.text.strip().startswith(('{%p ', '{%tr ')):
            para.paragraph_format.space_before = Pt(0)
            para.paragraph_format.space_after = Pt(0)
            for run in para.runs: run.font.size = Pt(2)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    d.save(OUT)
    print(OUT)


if __name__ == '__main__':
    create()
